"""Source-file ingestion: format detection, column normalisation and Parquet caching."""

from __future__ import annotations

import hashlib
import json
import logging
import re
from pathlib import Path

import pandas as pd

logger = logging.getLogger(__name__)

EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xls"}
CSV_SUFFIXES = {".csv", ".gz", ".zip"}
PARQUET_SUFFIXES = {".parquet", ".pq"}


class UnsupportedFormatError(ValueError):
    """Raised when a source file has an unsupported extension."""


def normalize_column_name(name: object) -> str:
    """Lower-case, strip and snake-case a column name (``"Loan Amnt "`` -> ``"loan_amnt"``)."""
    text = str(name).strip().lower()
    text = re.sub(r"[^0-9a-z]+", "_", text)
    return text.strip("_")


def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Return ``df`` with normalised column names; fails on collisions."""
    new_cols = [normalize_column_name(c) for c in df.columns]
    duplicates = sorted({c for c in new_cols if new_cols.count(c) > 1})
    if duplicates:
        raise ValueError(f"Column-name normalisation produced duplicates: {duplicates}")
    return df.set_axis(new_cols, axis=1)


def detect_format(path: Path) -> str:
    suffixes = [s.lower() for s in path.suffixes]
    if not suffixes:
        raise UnsupportedFormatError(f"Cannot detect format of {path} (no extension).")
    last = suffixes[-1]
    if last in EXCEL_SUFFIXES:
        return "excel"
    if last in PARQUET_SUFFIXES:
        return "parquet"
    if last == ".csv" or (last in CSV_SUFFIXES and ".csv" in suffixes):
        return "csv"
    raise UnsupportedFormatError(
        f"Unsupported file type '{''.join(suffixes)}' for {path}. Use CSV, Excel or Parquet."
    )


def list_excel_sheets(path: Path) -> list[str]:
    """Return sheet names without loading the workbook contents."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _stringify_object_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Make mixed-type object columns Parquet-safe (Excel cells may mix numbers and text)."""
    for col in df.columns:
        if df[col].dtype == object:
            series = df[col]
            df[col] = series.where(series.isna(), series.astype(str)).astype("string")
    return df


def read_source(path: str | Path, sheet: str | int | None = 0) -> pd.DataFrame:
    """Read a CSV, Excel or Parquet file and normalise its column names."""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(
            f"Source file not found: {path}. Place the Lending Club file under data/raw/ "
            "(see README 'Data placement')."
        )
    fmt = detect_format(path)
    if path.stat().st_size < 1024 and fmt == "excel":
        head = path.read_bytes()[:100]
        if head.startswith(b"version https://git-lfs"):
            raise ValueError(
                f"{path} is a Git LFS pointer, not the real workbook. Run `git lfs pull` "
                "or download the file from Figshare (see README)."
            )
    logger.info("Reading %s source %s", fmt, path)
    if fmt == "excel":
        df = pd.read_excel(path, sheet_name=0 if sheet is None else sheet, engine="openpyxl")
    elif fmt == "csv":
        df = pd.read_csv(path, low_memory=False)
    else:
        df = pd.read_parquet(path)
    df = normalize_columns(df)
    df = df.dropna(axis=0, how="all")
    df = _stringify_object_columns(df)
    logger.info("Loaded %s rows x %s columns", f"{len(df):,}", df.shape[1])
    return df


def file_digest(path: str | Path, algorithm: str = "sha256", chunk_size: int = 1 << 20) -> str:
    """Compute a hex digest of a file in streaming fashion."""
    h = hashlib.new(algorithm)
    with Path(path).open("rb") as fh:
        while chunk := fh.read(chunk_size):
            h.update(chunk)
    return h.hexdigest()


def _cache_key(path: Path, sheet: str | int | None) -> dict[str, object]:
    stat = path.stat()
    return {
        "source": str(path.name),
        "size": stat.st_size,
        "mtime_ns": stat.st_mtime_ns,
        "sheet": sheet,
    }


def load_raw(
    path: str | Path,
    cache_dir: str | Path | None = None,
    sheet: str | int | None = 0,
    use_cache: bool = True,
) -> pd.DataFrame:
    """Load the raw source, using a Parquet cache keyed by file size/mtime/sheet.

    Parsing the ~130 MB Figshare workbook takes several minutes; the cache makes
    subsequent runs take seconds. The original source file is never modified.
    """
    path = Path(path)
    if cache_dir is None or not use_cache:
        return read_source(path, sheet)
    cache_dir = Path(cache_dir)
    cache_file = cache_dir / "raw_cache.parquet"
    meta_file = cache_dir / "raw_cache.meta.json"
    if not path.is_file():
        raise FileNotFoundError(
            f"Source file not found: {path}. Place the Lending Club file under data/raw/."
        )
    key = _cache_key(path, sheet)
    if cache_file.is_file() and meta_file.is_file():
        try:
            cached_key = json.loads(meta_file.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            cached_key = None
        if cached_key == key:
            logger.info("Using cached raw data %s", cache_file)
            return pd.read_parquet(cache_file)
        logger.info("Raw cache is stale; re-reading source.")
    df = read_source(path, sheet)
    cache_dir.mkdir(parents=True, exist_ok=True)
    df.to_parquet(cache_file, index=False)
    meta_file.write_text(json.dumps(key, indent=2), encoding="utf-8")
    logger.info("Wrote raw cache %s", cache_file)
    return df
